"""Гараар оруулах үзүүлэлт (mode: manual).

  python ri.py manual template --activity ACT_01 [--date 2026-09-30]   → output/manual/*.xlsx
  (хариуцагч VALUE / VALUE_TEXT / SCORE / COMMENT-ийг бөглөнө)
  python ri.py manual load output/manual/manual_ACT_01_20260930.xlsx --env dev

SCORE бөглөсөн бол түүнийг (0 = Үнэлэхгүй), үгүй бол VALUE/VALUE_TEXT-ээс шатлалаар оноо тооцно.
Мөр бүр ri_result-д row_type=INDICATOR, entity_level=ACTIVITY, source_kind=MANUAL болж бичигдэнэ,
дараа нь үйл ажиллагааны нийт оноо дахин тооцогдоно.
"""
import csv
import datetime as dt
import getpass
import os
import uuid
from pathlib import Path

from . import store
from .config import OUTPUT_DIR, EnvConfig
from .db import connect
from .metadata import Catalog
from .runner import print_activity_scores, refresh_activities, score_value

COLUMNS = ["INDICATOR_ID", "BRANCH_ID", "ACTIVITY_ID", "ACTIVITY_NAME", "INDICATOR_NAME", "WEIGHT", "SCALE",
           "REQUIRED_INFO", "METHOD", "REPORT_DATE", "VALUE", "VALUE_TEXT", "SCORE", "COMMENT"]
INPUT_COLUMNS = ("VALUE", "VALUE_TEXT", "SCORE", "COMMENT")


def manual_indicators(cat: Catalog, activity_ids=None, include_all=False):
    acts = {cat.resolve_activity(a).activity_id for a in activity_ids or []}
    return [i for i in sorted(cat.indicators.values(), key=lambda i: i.indicator_id)
            if (i.mode == "manual" or include_all) and not i.example and i.enabled
            and (not acts or i.activity_id in acts)]


def write_template(cat: Catalog, activity_ids, report_date: dt.date, path: Path | None = None) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    inds = manual_indicators(cat, activity_ids)
    if not inds:
        raise ValueError("Гараар оруулах (mode: manual) үзүүлэлт алга")
    tag = "_".join(sorted({cat.activities[i.activity_id].branch_id or i.activity_id for i in inds})) \
        if activity_ids else "ALL"
    path = path or OUTPUT_DIR / "manual" / f"manual_{tag}_{report_date:%Y%m%d}.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "manual"
    ws.append(COLUMNS)
    for i in inds:
        act = cat.activities[i.activity_id]
        ws.append([i.indicator_id, act.branch_id, i.activity_id, act.name_mn, i.name_mn, i.weight, i.scale_text,
                   i.master.get("required_info", ""), i.master.get("method", ""), report_date.isoformat(),
                   None, None, None, None])
    fill = PatternFill("solid", fgColor="FFF2CC")
    for c, name in enumerate(COLUMNS, 1):
        cell = ws.cell(1, c)
        cell.font = Font(bold=True)
        if name in INPUT_COLUMNS:
            cell.fill = fill
            for r in range(2, ws.max_row + 1):
                ws.cell(r, c).fill = fill
    widths = {"INDICATOR_NAME": 45, "SCALE": 40, "REQUIRED_INFO": 35, "METHOD": 35, "ACTIVITY_NAME": 30,
              "COMMENT": 30, "VALUE_TEXT": 20}
    for c, name in enumerate(COLUMNS, 1):
        ws.column_dimensions[ws.cell(1, c).column_letter].width = widths.get(name, 13)
        if name in ("SCALE", "REQUIRED_INFO", "METHOD", "INDICATOR_NAME"):
            for r in range(2, ws.max_row + 1):
                ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "F2"
    wb.save(path)
    return path


def read_file(path: Path) -> list[dict]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True, data_only=True).active
        it = ws.iter_rows(values_only=True)
        head = [str(h or "").strip().upper() for h in next(it)]
        return [dict(zip(head, r)) for r in it if any(v not in (None, "") for v in r)]
    with path.open(encoding="utf-8-sig", newline="") as f:
        return [{(k or "").strip().upper(): v for k, v in r.items()} for r in csv.DictReader(f)]


def _num(v):
    if v in (None, ""):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "").replace(",", ".")
    pct = s.endswith("%")
    try:
        x = float(s.rstrip("%"))
    except ValueError:
        return None
    return x / 100 if pct else x


def load(cfg: EnvConfig, cat: Catalog, path: Path, default_date: dt.date | None = None,
         dry_run: bool = False) -> int:
    rows = read_file(path)
    run_id = str(uuid.uuid4())
    user = os.environ.get("RI_USER") or getpass.getuser()
    errors, prepared = [], []
    for n, r in enumerate(rows, 2):
        iid = str(r.get("INDICATOR_ID") or "").strip().upper()
        ind = cat.indicators.get(iid)
        if ind is None:
            errors.append(f"мөр {n}: үзүүлэлт олдсонгүй '{iid}'")
            continue
        if ind.mode != "manual":
            errors.append(f"мөр {n}: {iid} нь mode={ind.mode} (SQL-ээр тооцогдоно) — manual биш")
            continue
        value, text, score_in = _num(r.get("VALUE")), r.get("VALUE_TEXT"), r.get("SCORE")
        text = None if text in (None, "") else str(text).strip()
        if value is None and text is None and score_in in (None, ""):
            continue                                    # бөглөөгүй — алгасна
        try:
            d = dt.date.fromisoformat(str(r.get("REPORT_DATE"))[:10]) if r.get("REPORT_DATE") else default_date
        except ValueError:
            errors.append(f"мөр {n}: REPORT_DATE буруу '{r.get('REPORT_DATE')}'")
            continue
        if d is None:
            errors.append(f"мөр {n}: REPORT_DATE хоосон (--date өгнө үү)")
            continue
        if score_in not in (None, ""):
            s = _num(score_in)
            if s is None or s not in (0, 1, 2, 3, 4, 5):
                errors.append(f"мөр {n}: SCORE 0–5 байх ёстой ('{score_in}')")
                continue
            score, label = (None, "Үнэлэхгүй") if s == 0 else (s, "гараар")
        else:
            score, label, _, _ = score_value(ind, value, text)
        act = cat.activities[ind.activity_id]
        if not act.branch_id:
            errors.append(f"мөр {n}: {iid} — {act.activity_id}-д BRANCHID алга (config/activity_map.csv)")
            continue
        prepared.append({
            "report_date": d, "period_start": None, "row_type": "INDICATOR", "group_no": act.group_no,
            "activity_id": act.activity_id, "branch_id": str(act.branch_id or ""), "activity_name": act.name_mn,
            "indicator_id": iid, "indicator_name": ind.name_mn, "calc_id": iid, "source_kind": "MANUAL",
            "entity_level": "ACTIVITY", "entity_id": str(act.branch_id or act.activity_id),
            "entity_name": act.name_mn, "metric": ind.score_metric, "value": value, "value_text": text,
            "score": score, "score_label": label, "weight": ind.weight, "data_status": "OK",
            "score_type": ind.score_type, "scale_version": ind.scale_version,
            "comment": None if r.get("COMMENT") in (None, "") else str(r.get("COMMENT")),
            "entered_by": user, "run_id": run_id, "sql_version": f"{ind.version}+manual"})

    for e in errors:
        print(f"ERROR {e}")
    print(f"{path.name}: {len(prepared)} мөр бэлэн, {len(errors)} алдаа")
    for p in prepared[:40] if len(prepared) > 60 else prepared:
        sc = "—" if p["score"] is None else f"{p['score']:g}"
        v = p["value_text"] if p["value"] is None else p["value"]
        print(f"  {p['indicator_id']} {p['report_date']}  утга={v}  оноо={sc} ({p['score_label']})")
    if errors or dry_run or not prepared:
        if errors:
            print("Алдааг засаад дахин ачаална уу (юу ч бичигдсэнгүй).")
        return 1 if errors else 0

    from .runner import sync_meta
    sync_meta(cfg, cat)
    touched: dict[dt.date, set] = {}
    keys = [(p["indicator_id"], p["report_date"]) for p in prepared]
    if len(set(keys)) != len(keys):
        print("ERROR файлд ижил (INDICATOR_ID, REPORT_DATE) давхардсан — юу ч бичигдсэнгүй")
        return 1
    with connect(cfg, cfg.result_connection) as conn:
        store.ensure_base(cfg, conn)
        # Бүгдийг НЭГ insert → ClickHouse-оос буцааж уншиж мөр бүрийг тулгана (зөрвөл энэ run-ийг устгана)
        store.insert_result_verified(cfg, conn, prepared, {"row_type": "INDICATOR", "source_kind": "MANUAL"}, run_id)
        replaced = store.keep_only_run_many(cfg, conn, keys, run_id)
        now = dt.datetime.now()
        store.write_audits(cfg, conn, [{
            "run_id": run_id, "run_mode": "MANUAL", "activity_id": p["activity_id"],
            "indicator_id": p["indicator_id"], "report_date": p["report_date"],
            "execution_status": "SUCCESS", "started_at": now, "finished_at": now, "duration_seconds": 0,
            "result_row_count": 1, "affected_rows": None, "executed_by": user, "environment": cfg.env,
            "sql_version": p["sql_version"], "connection_name": "manual", "target_table": store.RESULT_TABLE,
            "validation_detail": f"file={path.name}"} for p in prepared])
        for p in prepared:
            touched.setdefault(p["report_date"], set()).add(p["activity_id"])
    print(f"{len(prepared)} мөр бичиж буцааж уншиж тулгав, өмнөх {replaced} мөр ri_history-д")
    acts = []
    for d, ids in sorted(touched.items()):
        acts += refresh_activities(cfg, cat, ids, d, run_id)
    print(f"\n{cfg.result_db}.{store.RESULT_TABLE}-д бичлээ (run_id={run_id})")
    print_activity_scores(acts)
    return 0
